import csv
import io
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.domain import (
    Brand,
    Category,
    Characteristic,
    LexicalVariant,
    ProductBrand,
    ProductService,
    Provider,
    ProviderProductService,
    RequirementTemplate,
    TemplateField,
    UnitMeasure,
)

SUPPORTED_TYPES = {
    "categorias",
    "productos",
    "variantes",
    "marcas",
    "caracteristicas",
    "unidades",
    "proveedores",
    "proveedor_productos",
    "producto_marcas",
    "plantillas",
    "plantilla_campos",
}


class CatalogImportService:
    def __init__(self, db: Session):
        self.db = db

    def import_bytes(self, filename: str, content: bytes, entity_type: str) -> int:
        if entity_type not in SUPPORTED_TYPES:
            raise ValueError(f"Tipo no soportado: {entity_type}")
        rows = self._rows(filename, content)
        count = 0
        for row in rows:
            if self._import_row(entity_type, row):
                count += 1
        self.db.commit()
        return count

    def _rows(self, filename: str, content: bytes) -> list[dict]:
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv":
            return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        if suffix == ".xlsx":
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            values = list(ws.iter_rows(values_only=True))
            if not values:
                return []
            headers = [str(v).strip() if v is not None else "" for v in values[0]]
            return [dict(zip(headers, row)) for row in values[1:]]
        raise ValueError("Solo se permiten archivos CSV o XLSX")

    @staticmethod
    def _value(row: dict, *names: str) -> str:
        for name in names:
            value = row.get(name)
            if value is not None and str(value).strip():
                return str(value).strip()
        return ""

    @staticmethod
    def _bool(value: str, default: bool = False) -> bool:
        if value == "":
            return default
        return value.lower() in {"1", "true", "si", "sí", "yes", "y"}

    def _product(self, name: str) -> ProductService | None:
        return self.db.scalar(select(ProductService).where(ProductService.name.ilike(name)))

    def _brand(self, name: str) -> Brand | None:
        return self.db.scalar(select(Brand).where(Brand.name.ilike(name)))

    def _provider(self, name: str) -> Provider | None:
        return self.db.scalar(select(Provider).where(Provider.name.ilike(name)))

    def _characteristic(self, key_or_name: str) -> Characteristic | None:
        return self.db.scalar(
            select(Characteristic).where(
                (Characteristic.key.ilike(key_or_name)) | (Characteristic.name.ilike(key_or_name))
            )
        )

    def _import_row(self, entity_type: str, row: dict) -> bool:
        if entity_type == "categorias":
            name = self._value(row, "name", "nombre")
            if not name or self.db.scalar(select(Category).where(Category.name.ilike(name))):
                return False
            self.db.add(Category(name=name)); return True

        if entity_type == "productos":
            name = self._value(row, "name", "nombre")
            category_name = self._value(row, "category", "categoria")
            kind = self._value(row, "kind", "tipo") or "PRODUCTO"
            if not name or self._product(name):
                return False
            category = self.db.scalar(select(Category).where(Category.name.ilike(category_name))) if category_name else None
            if category_name and not category:
                raise ValueError(f"Categoría no encontrada: {category_name}")
            self.db.add(ProductService(name=name, category_id=category.id if category else None, kind=kind.upper(), active=True)); return True

        if entity_type == "variantes":
            variant = self._value(row, "variant", "variante")
            normalized = self._value(row, "normalized_value", "normalizado")
            product_name = self._value(row, "product", "producto")
            requires_confirmation = self._bool(self._value(row, "requires_confirmation", "requiere_confirmacion"))
            if not variant or not normalized or self.db.scalar(select(LexicalVariant).where(LexicalVariant.variant.ilike(variant))):
                return False
            product = self._product(product_name) if product_name else None
            if product_name and not product:
                raise ValueError(f"Producto no encontrado: {product_name}")
            self.db.add(LexicalVariant(
                variant=variant,
                normalized_value=normalized,
                product_service_id=product.id if product else None,
                requires_confirmation=requires_confirmation,
                transformation_type="VARIANTE_LEXICA",
            )); return True

        if entity_type == "marcas":
            name = self._value(row, "name", "nombre")
            if not name or self._brand(name):
                return False
            self.db.add(Brand(name=name)); return True

        if entity_type == "caracteristicas":
            key = self._value(row, "key", "clave")
            name = self._value(row, "name", "nombre")
            if not key or not name or self.db.scalar(select(Characteristic).where(Characteristic.key.ilike(key))):
                return False
            self.db.add(Characteristic(key=key, name=name)); return True

        if entity_type == "unidades":
            name = self._value(row, "name", "nombre")
            symbol = self._value(row, "symbol", "simbolo")
            if not name or not symbol or self.db.scalar(select(UnitMeasure).where(UnitMeasure.symbol.ilike(symbol))):
                return False
            self.db.add(UnitMeasure(name=name, symbol=symbol)); return True

        if entity_type == "proveedores":
            name = self._value(row, "name", "nombre")
            email = self._value(row, "email", "correo")
            if not name or not email or self._provider(name):
                return False
            self.db.add(Provider(name=name, email=email, active=True)); return True

        if entity_type == "proveedor_productos":
            provider_name = self._value(row, "provider", "proveedor")
            product_name = self._value(row, "product", "producto")
            provider, product = self._provider(provider_name), self._product(product_name)
            if not provider or not product:
                raise ValueError(f"Relación inválida proveedor/producto: {provider_name} / {product_name}")
            exists = self.db.get(ProviderProductService, (provider.id, product.id))
            if exists: return False
            self.db.add(ProviderProductService(provider_id=provider.id, product_service_id=product.id)); return True

        if entity_type == "producto_marcas":
            product_name = self._value(row, "product", "producto")
            brand_name = self._value(row, "brand", "marca")
            product, brand = self._product(product_name), self._brand(brand_name)
            if not product or not brand:
                raise ValueError(f"Relación inválida producto/marca: {product_name} / {brand_name}")
            if self.db.get(ProductBrand, (product.id, brand.id)): return False
            self.db.add(ProductBrand(product_service_id=product.id, brand_id=brand.id)); return True

        if entity_type == "plantillas":
            product_name = self._value(row, "product", "producto")
            name = self._value(row, "name", "nombre")
            product = self._product(product_name)
            if not product or not name:
                raise ValueError("Plantilla requiere columnas producto y nombre")
            exists = self.db.scalar(select(RequirementTemplate).where(RequirementTemplate.product_service_id == product.id, RequirementTemplate.name.ilike(name)))
            if exists: return False
            self.db.add(RequirementTemplate(product_service_id=product.id, name=name, active=True)); return True

        if entity_type == "plantilla_campos":
            template_name = self._value(row, "template", "plantilla")
            characteristic_key = self._value(row, "characteristic", "caracteristica")
            question = self._value(row, "question", "pregunta")
            required = self._bool(self._value(row, "required", "obligatorio"), default=True)
            template = self.db.scalar(select(RequirementTemplate).where(RequirementTemplate.name.ilike(template_name)))
            characteristic = self._characteristic(characteristic_key)
            if not template or not characteristic or not question:
                raise ValueError("Campo de plantilla requiere plantilla, característica y pregunta válidas")
            exists = self.db.scalar(select(TemplateField).where(TemplateField.template_id == template.id, TemplateField.characteristic_id == characteristic.id))
            if exists: return False
            self.db.add(TemplateField(template_id=template.id, characteristic_id=characteristic.id, required=required, question_text=question)); return True

        return False
